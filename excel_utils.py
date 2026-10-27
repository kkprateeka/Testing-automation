from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill



def create_excel(results, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Report"

    # -------------------------
    # ✅ HEADER
    # -------------------------

    headers = [
        "#",
        "Test Name",
        "Group",
        "Step Type",
        "Element",
        "Severity",
        "Status",
        "Execution Time (s)",
        "Screenshot",
        "Message",
        "AI Analysis"
    ]
    ws.append(headers)

    # style header
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # -------------------------
    # ✅ DATA
    # -------------------------
    row_num = 2  # start after header
    for i, r in enumerate(results, start=1):
        steps = r.get("steps", [])
        # 🔹 If no step-level data → fallback
        if not steps:
            ws.append([
                i,
                r.get("name", ""),
                "General",
                "N/A",
                "N/A",
                "Medium",
                r.get("status", ""),
                round(r.get("duration", 0), 2),
                "",
                r.get("details", "")
            ])
            # add screenshot hyperlink
            if r.get("screenshot"):
                cell = ws.cell(row=row_num, column=9)
                cell.value = "View"
                cell.hyperlink = r["screenshot"]
                cell.style = "Hyperlink"
                
            row_num += 1
            continue
        # 🔥 STEP LEVEL ROWS
        for step in steps:
            # 🔥 GROUP
            group = "General"
            if step.get("type") == "upload":
                group = "Upload"
            elif step.get("type") == "button":
                group = "Navigation"
            elif step.get("type") == "input":
                group = "Form"
                
            # 🔥 SEVERITY
            severity = "Medium"
            if step.get("status") == "FAIL":
                severity = "Critical"
            ws.append([
                i,
                r.get("name", ""),
                group,
                step.get("type", ""),
                step.get("name", ""),
                severity,
                step.get("status", ""),
                round(r.get("duration", 0), 2),
                "",
                step.get("message", ""),
                r.get("ai_analysis", "")
            ])
            # ✅ add hyperlink AFTER append
            screenshot_path = step.get("screenshot", "")
            print("SCREENSHOT LINK:", screenshot_path)

            if screenshot_path:
                cell = ws.cell(row=row_num, column=9)
                cell.value = "View"
                cell.hyperlink = screenshot_path
                cell.style = "Hyperlink"
            row_num += 1
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        status_cell = row[6]  # Status column
        
        if status_cell.value == "PASS":
            status_cell.fill = green
        elif status_cell.value == "FAIL":
            status_cell.fill = red

    # -------------------------
    # ✅ AUTO WIDTH
    # -------------------------
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass

        ws.column_dimensions[col_letter].width = max_length + 2

    # -------------------------
    # 📊 SUMMARY SHEET
    # -------------------------
     
    summary = wb.create_sheet("Summary")
    total = 0
    passed = 0
    failed = 0
    
    for r in results:
        for step in r.get("steps", []):
            total += 1
            if step.get("status") == "PASS":
                passed += 1
            else:
                failed += 1
    pass_percent = 0
    if total > 0:
        pass_percent = round((passed / total) * 100, 2)
        
    summary.append(["Metric", "Value"])
    summary.append(["Total Checks", total])
    summary.append(["Passed", passed])
    summary.append(["Failed", failed])
    summary.append(["Pass Percentage", f"{pass_percent}%"])
    
    # style summary
    for cell in summary[1]:
        cell.font = Font(bold=True)
        
    # -------------------------
    # ✅ SAVE
    # -------------------------
    wb.save(path)